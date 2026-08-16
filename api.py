from pathlib import Path
from typing import Dict, Any
from numbers import Number
from uuid import uuid4
from datetime import datetime, timezone
import os
from urllib.parse import quote_plus

from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import JSONResponse
from fastapi.responses import FileResponse
from fastapi.responses import RedirectResponse
from fastapi.staticfiles import StaticFiles

from config import (
    TEMPLATE_CATEGORY_ROW_RANGES,
    TEMPLATE_EXPENSE_CATEGORY_ROW_RANGES,
    TEMPLATE_INCOME_CATEGORY_ROW_RANGES,
    TEMPLATE_INVESTMENT_CATEGORY_ROW_RANGES,
    get_paths,
    get_state,
    load_env,
    update_state,
)
from pipeline import run_for_month, prepare_for_month, finalize_for_month
from data import monthToExpenseColDict
from openpyxl import load_workbook
from handleExcel import (
    parse_ai_output_lines,
    get_allowed_categories,
    get_allowed_category_groups,
    build_ai_output_from_items,
    build_review_items,
)
from google_drive_service import (
    create_oauth_start_url,
    disconnect_drive,
    get_drive_status,
    handle_oauth_callback,
    upload_month_files,
)


app = FastAPI(title="AutoExpenses API")
_PENDING_REVIEWS: Dict[str, Dict[str, Any]] = {}
_PIPELINE_PROGRESS: Dict[str, Dict[str, Any]] = {}

# Ensure API routes (including Drive status/connect) can read .env configuration.
load_env()


BASE_DIR = Path(__file__).resolve().parent
FRONTEND_DIR = BASE_DIR / "frontend"
STATIC_DIR = FRONTEND_DIR / "static"
NEXT_OUT_DIR = BASE_DIR / "frontend-next" / "out"
NEXT_ASSETS_DIR = NEXT_OUT_DIR / "_next"

# Serve the web UI (index.html) + its static assets from /static/*
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
if NEXT_ASSETS_DIR.exists():
    app.mount("/_next", StaticFiles(directory=str(NEXT_ASSETS_DIR)), name="next-assets")


def _user_friendly_pipeline_error_message(error: Exception) -> str:
    """Map internal pipeline failures to short user-facing messages."""
    raw_message = str(error)
    lower_message = raw_message.lower()

    if "step 1 (download max)" in lower_message:
        return (
            "Failed to download the Max Excel file. "
            "Please check that the Max site is available and try again."
        )
    if "step 2 (leumi download/merge)" in lower_message:
        return (
            "Failed while downloading or merging Leumi data. "
            "Please check your Leumi access and try again."
        )
    if "reconciliation" in lower_message:
        return (
            "Downloaded data failed reconciliation checks. "
            "Please review the exports and try again."
        )
    if "step 3 (getexpenses/categorize)" in lower_message:
        return (
            "Failed while categorizing expenses. "
            "Please verify your API configuration and try again."
        )
    if "step 5 (fillcells)" in lower_message:
        return (
            "Failed while writing results to the workbook. "
            "Please ensure the file is not open and try again."
        )
    return "The monthly run failed. Please try again."


@app.exception_handler(Exception)
async def unhandled_exception_handler(request: Request, exc: Exception) -> JSONResponse:
    # Ensure the frontend always gets JSON, even for unexpected server errors.
    return JSONResponse(
        status_code=500,
        content={
            "detail": "Something went wrong on the server. Please try again.",
            "error_code": "UNEXPECTED_SERVER_ERROR",
        },
    )


@app.get("/")
def index() -> FileResponse:
    next_index = NEXT_OUT_DIR / "index.html"
    if next_index.exists():
        return FileResponse(str(next_index))
    return FileResponse(str(FRONTEND_DIR / "index.html"))


@app.post("/run-month")
def run_month(payload: Dict[str, Any]) -> Dict[str, Any]:
    year = str(payload.get("year"))
    month = str(payload.get("month"))
    include_leumi = bool(payload.get("include_leumi", False))

    if not year.isdigit() or not month.isdigit():
        raise HTTPException(
            status_code=400,
            detail="Invalid year/month. Please choose valid numeric values.",
        )

    month_int = int(month)
    if month_int < 1 or month_int > 12:
        raise HTTPException(
            status_code=400,
            detail="Invalid month. Please choose a month between 1 and 12.",
        )

    try:
        src, out = run_for_month(year, month, include_leumi=include_leumi)
    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=_user_friendly_pipeline_error_message(error),
        ) from error

    return {
        "source_excel": str(src),
        "output_excel": str(out),
        "state": get_state(),
    }


@app.post("/run-month/prepare")
def run_month_prepare(payload: Dict[str, Any]) -> Dict[str, Any]:
    year = str(payload.get("year"))
    month = str(payload.get("month"))
    include_leumi = bool(payload.get("include_leumi", False))

    if not year.isdigit() or not month.isdigit():
        raise HTTPException(
            status_code=400,
            detail="Invalid year/month. Please choose valid numeric values.",
        )

    month_int = int(month)
    if month_int < 1 or month_int > 12:
        raise HTTPException(
            status_code=400,
            detail="Invalid month. Please choose a month between 1 and 12.",
        )

    run_token = str(payload.get("run_token", "")).strip() or str(uuid4())

    def _update_progress(step_text: str) -> None:
        _PIPELINE_PROGRESS[run_token] = {
            "step": step_text,
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "done": False,
        }

    _update_progress("Starting monthly run...")

    try:
        src, out, ai_output, source_expenses_dict, capture_warnings = prepare_for_month(
            year=year,
            month=month,
            include_leumi=include_leumi,
            progress_callback=_update_progress,
        )
    except Exception as error:
        _PIPELINE_PROGRESS[run_token] = {
            "step": "Run failed.",
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "done": True,
        }
        raise HTTPException(
            status_code=500,
            detail=_user_friendly_pipeline_error_message(error),
        ) from error

    parsed_items, parse_errors = parse_ai_output_lines(ai_output)
    allowed_category_groups = get_allowed_category_groups(str(out))
    allowed_categories = [
        category
        for group in allowed_category_groups
        for category in group["categories"]
    ]
    allow_duplicate_exclusion = True
    if isinstance(source_expenses_dict, list) and source_expenses_dict:
        first = source_expenses_dict[0]
        if isinstance(first, dict) and "_card_export_healthy" in first:
            allow_duplicate_exclusion = bool(first.get("_card_export_healthy"))
    review_items, review_warnings = build_review_items(
        parsed_items,
        parse_errors,
        expenses_dict=source_expenses_dict,
        allow_duplicate_exclusion=allow_duplicate_exclusion,
    )
    warnings = [*capture_warnings, *review_warnings]

    review_token = str(uuid4())
    _PENDING_REVIEWS[review_token] = {
        "year": year,
        "month": month,
        "output_workbook_path": str(out),
        "expenses_dict": source_expenses_dict,
    }

    return {
        "run_token": run_token,
        "review_token": review_token,
        "source_excel": str(src),
        "output_excel": str(out),
        "allowed_categories": allowed_categories,
        "category_groups": allowed_category_groups,
        "items": review_items,
        "parse_errors": parse_errors,
        "warnings": warnings,
    }


@app.post("/run-month/finalize")
def run_month_finalize(payload: Dict[str, Any]) -> Dict[str, Any]:
    review_token = str(payload.get("review_token", "")).strip()
    reviewed_items = payload.get("reviewed_items", [])
    if not review_token:
        raise HTTPException(status_code=400, detail="Missing review token.")

    review_state = _PENDING_REVIEWS.get(review_token)
    if review_state is None:
        raise HTTPException(
            status_code=400,
            detail="Review session not found or expired. Please run prepare again.",
        )

    if not isinstance(reviewed_items, list):
        raise HTTPException(status_code=400, detail="reviewed_items must be a list.")

    categorized_output = build_ai_output_from_items(reviewed_items)
    run_token = str(payload.get("run_token", "")).strip()

    def _update_progress(step_text: str) -> None:
        if not run_token:
            return
        _PIPELINE_PROGRESS[run_token] = {
            "step": step_text,
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "done": False,
        }

    if run_token:
        _update_progress("Step 5/5: Filling workbook...")

    try:
        final_path, yearly_total_path = finalize_for_month(
            year=review_state["year"],
            month=review_state["month"],
            output_workbook_path=Path(review_state["output_workbook_path"]),
            categorized_output=categorized_output,
            progress_callback=_update_progress if run_token else None,
            expenses_dict=review_state.get("expenses_dict"),
        )
    except Exception as error:
        if run_token:
            _PIPELINE_PROGRESS[run_token] = {
                "step": "Run failed.",
                "updated_at": datetime.now(timezone.utc).isoformat(),
                "done": True,
            }
        raise HTTPException(
            status_code=500,
            detail=_user_friendly_pipeline_error_message(error),
        ) from error
    finally:
        _PENDING_REVIEWS.pop(review_token, None)

    if run_token:
        _PIPELINE_PROGRESS[run_token] = {
            "step": "Done.",
            "updated_at": datetime.now(timezone.utc).isoformat(),
            "done": True,
        }

    drive_status = get_drive_status()
    drive_upload: Dict[str, Any] = {
        "attempted": False,
        "connected": bool(drive_status.get("connected", False)),
        "success": False,
        "error": None,
        "folder_path": None,
        "monthly_file": None,
        "yearly_file": None,
    }
    if drive_status.get("configured") and drive_status.get("connected"):
        drive_upload["attempted"] = True
        try:
            upload_result = upload_month_files(
                year=str(review_state["year"]),
                month=str(review_state["month"]),
                monthly_path=final_path,
                yearly_total_path=yearly_total_path,
            )
            drive_upload.update(upload_result)
            drive_upload["success"] = True
        except Exception as upload_error:
            drive_upload["success"] = False
            drive_upload["error"] = str(upload_error)

    return {
        "output_excel": str(final_path),
        "yearly_output_excel": str(yearly_total_path),
        "drive_upload": drive_upload,
        "state": get_state(),
    }


@app.get("/run-month/progress/{run_token}")
def run_month_progress(run_token: str) -> Dict[str, Any]:
    progress = _PIPELINE_PROGRESS.get(run_token)
    if progress is None:
        return {"step": "", "updated_at": None, "done": False}
    return progress


@app.get("/state")
def read_state() -> Dict[str, Any]:
    return get_state()


@app.get("/drive/status")
def drive_status() -> Dict[str, Any]:
    return get_drive_status()


@app.post("/drive/connect/start")
def drive_connect_start() -> Dict[str, str]:
    try:
        return create_oauth_start_url()
    except Exception as error:
        raise HTTPException(status_code=400, detail=str(error)) from error


@app.get("/drive/connect/callback")
def drive_connect_callback(code: str = "", state: str = "") -> RedirectResponse:
    if not code or not state:
        raise HTTPException(status_code=400, detail="Missing OAuth callback code/state.")

    frontend_url = os.environ.get("DASHBOARD_BASE_URL", "http://localhost:3000").rstrip("/")
    try:
        handle_oauth_callback(code=code, state=state)
        return RedirectResponse(
            url=f"{frontend_url}/?drive_connected=1",
            status_code=302,
        )
    except Exception as error:
        error_text = quote_plus(str(error))
        return RedirectResponse(
            url=f"{frontend_url}/?drive_connected=0&drive_error={error_text}",
            status_code=302,
        )


@app.post("/drive/disconnect")
def drive_disconnect() -> Dict[str, Any]:
    return disconnect_drive()


@app.post("/assets")
def write_assets(payload: Dict[str, Any]) -> Dict[str, Any]:
    cars = payload.get("cars", [])
    investments = payload.get("investments", [])
    update_state({"cars": cars, "investments": investments})
    return get_state()


@app.get("/assets")
def read_assets() -> Dict[str, Any]:
    state = get_state()
    return {
        "cars": state.get("cars", []),
        "investments": state.get("investments", []),
        "leumi_balance": state.get("leumi_balance"),
    }


@app.get("/expenses/summary")
def expenses_summary(year: int, month: int) -> Dict[str, float]:
    """Summarize expenses by category for a given year/month."""
    return _summarize_sheet_row_ranges(
        year=year,
        month=month,
        row_ranges=list(TEMPLATE_EXPENSE_CATEGORY_ROW_RANGES),
    )


@app.get("/income/summary")
def income_summary(year: int, month: int) -> Dict[str, float]:
    """Summarize income by category for a given year/month."""
    return _summarize_sheet_row_ranges(
        year=year,
        month=month,
        row_ranges=list(TEMPLATE_INCOME_CATEGORY_ROW_RANGES),
    )


@app.get("/investments/summary")
def investments_summary(year: int, month: int) -> Dict[str, float]:
    """Summarize investment rows by category for a given year/month."""
    return _summarize_sheet_row_ranges(
        year=year,
        month=month,
        row_ranges=list(TEMPLATE_INVESTMENT_CATEGORY_ROW_RANGES),
    )


@app.get("/categories")
def list_categories() -> Dict[str, Any]:
    """Return template category groups for filters and pickers."""
    paths = get_paths()
    template = paths.template_expenses
    if not template.exists():
        return {"groups": [], "categories": []}
    groups = get_allowed_category_groups(str(template))
    categories = [category for group in groups for category in group["categories"]]
    return {"groups": groups, "categories": categories}


@app.get("/categories/timeline")
def category_timeline(
    category: str,
    mode: str = "year",
    year: int | None = None,
    trailing_months: int = 12,
    start_year: int | None = None,
    start_month: int | None = None,
    end_year: int | None = None,
    end_month: int | None = None,
) -> Dict[str, Any]:
    """
    Build a timeline of monthly amounts for a single category.
    Uses the same mode/range options as /totals/timeline.
    """
    selected_category = (category or "").strip()
    if not selected_category:
        raise HTTPException(status_code=400, detail='Query param "category" is required.')

    selected_months = _resolve_timeline_months(
        mode=mode,
        year=year,
        trailing_months=trailing_months,
        start_year=start_year,
        start_month=start_month,
        end_year=end_year,
        end_month=end_month,
    )
    points = [
        {
            "year": month_year,
            "month": month,
            "label": f"{month_year}-{str(month).zfill(2)}",
            "amount": _monthly_category_amount(month_year, month, selected_category),
        }
        for month_year, month in selected_months
    ]
    return {"mode": mode, "category": selected_category, "points": points}


@app.get("/totals/timeline")
def totals_timeline(
    mode: str = "year",
    year: int | None = None,
    trailing_months: int = 12,
    start_year: int | None = None,
    start_month: int | None = None,
    end_year: int | None = None,
    end_month: int | None = None,
) -> Dict[str, Any]:
    """
    Build a timeline of month totals for spending/income/investments.
    mode:
      - "year": use all 12 months from selected year.
      - "trailing": use latest N available output months.
      - "range": use inclusive start/end year-month.
    """
    selected_months = _resolve_timeline_months(
        mode=mode,
        year=year,
        trailing_months=trailing_months,
        start_year=start_year,
        start_month=start_month,
        end_year=end_year,
        end_month=end_month,
    )
    points = [
        {
            "year": month_year,
            "month": month,
            "label": f"{month_year}-{str(month).zfill(2)}",
            **_monthly_totals(month_year, month),
        }
        for month_year, month in selected_months
    ]
    return {"mode": mode, "points": points}


def _resolve_timeline_months(
    *,
    mode: str,
    year: int | None,
    trailing_months: int,
    start_year: int | None,
    start_month: int | None,
    end_year: int | None,
    end_month: int | None,
) -> list[tuple[int, int]]:
    """Resolve (year, month) points for timeline endpoints."""
    paths = get_paths()
    available_months = _available_output_months(paths.data_dir)
    if not available_months:
        return []

    if mode == "year":
        selected_year = year or datetime.now().year
        return [(selected_year, month) for month in range(1, 13)]

    if mode == "trailing":
        safe_trailing_months = max(1, min(120, int(trailing_months)))
        return sorted(available_months)[-safe_trailing_months:]

    if mode == "range":
        if None in (start_year, start_month, end_year, end_month):
            raise HTTPException(
                status_code=400,
                detail=(
                    'Range mode requires "start_year", "start_month", '
                    '"end_year", and "end_month".'
                ),
            )
        if not (1 <= int(start_month) <= 12 and 1 <= int(end_month) <= 12):
            raise HTTPException(
                status_code=400,
                detail="Range mode month values must be between 1 and 12.",
            )

        start = (int(start_year), int(start_month))
        end = (int(end_year), int(end_month))
        if start > end:
            raise HTTPException(
                status_code=400,
                detail="Range mode start must be earlier than or equal to end.",
            )
        return _month_span(start=start, end=end)

    raise HTTPException(
        status_code=400,
        detail='Invalid mode. Please use "year", "trailing", or "range".',
    )


def _summarize_sheet_row_ranges(
    *,
    year: int,
    month: int,
    row_ranges: list[tuple[int, int]],
) -> Dict[str, float]:
    paths = get_paths()

    month_padded = str(month).zfill(2)
    # pipeline writes per-month workbooks like: data/{year}/expenses_output_{year}_{MM}.xlsx
    candidate_path = (
        paths.data_dir / str(year) / f"expenses_output_{year}_{month_padded}.xlsx"
    )
    if not candidate_path.exists():
        # Fallbacks for older flows / manual runs.
        flat_data_path = paths.data_dir / f"expenses_output_{year}_{month_padded}.xlsx"
        if flat_data_path.exists():
            candidate_path = flat_data_path
        else:
            candidate_path = (
                paths.current_expenses if paths.current_expenses.exists() else None
            )
    if not candidate_path:
        return {}

    wb = load_workbook(str(candidate_path), data_only=True)
    ws = wb.active

    if month not in monthToExpenseColDict:
        return {}
    col_letter = monthToExpenseColDict[month]

    category_col = "B"
    summary: Dict[str, float] = {}
    for first_row, last_row_exclusive in row_ranges:
        for row in range(first_row, last_row_exclusive):
            category = ws[f"{category_col}{row}"].value
            if not category:
                continue
            cell = ws[f"{col_letter}{row}"]
            raw_value = cell.value
            if raw_value is None or raw_value == "":
                continue

            amount: float | None = None
            if isinstance(raw_value, Number):
                amount = float(raw_value)
            elif isinstance(raw_value, str):
                # Some template rows in month columns may contain labels (e.g., month names).
                normalized = raw_value.replace(",", "").strip()
                try:
                    amount = float(normalized)
                except ValueError:
                    continue

            if amount is not None:
                summary[str(category)] = summary.get(str(category), 0.0) + amount

    return summary


def _monthly_totals(year: int, month: int) -> Dict[str, float]:
    spending = sum(
        _summarize_sheet_row_ranges(
            year=year,
            month=month,
            row_ranges=list(TEMPLATE_EXPENSE_CATEGORY_ROW_RANGES),
        ).values()
    )
    income = sum(
        _summarize_sheet_row_ranges(
            year=year,
            month=month,
            row_ranges=list(TEMPLATE_INCOME_CATEGORY_ROW_RANGES),
        ).values()
    )
    investments = sum(
        _summarize_sheet_row_ranges(
            year=year,
            month=month,
            row_ranges=list(TEMPLATE_INVESTMENT_CATEGORY_ROW_RANGES),
        ).values()
    )
    return {
        "spending": float(spending),
        "income": float(income),
        "investments": float(investments),
    }


def _monthly_category_amount(year: int, month: int, category: str) -> float:
    summary = _summarize_sheet_row_ranges(
        year=year,
        month=month,
        row_ranges=list(TEMPLATE_CATEGORY_ROW_RANGES),
    )
    return float(summary.get(category, 0.0))


def _available_output_months(data_dir: Path) -> list[tuple[int, int]]:
    """
    Return all months that have generated output workbook files.
    Pattern: expenses_output_YYYY_MM.xlsx
    """
    months: set[tuple[int, int]] = set()
    for file_path in data_dir.glob("**/expenses_output_????_??.xlsx"):
        stem_parts = file_path.stem.split("_")
        if len(stem_parts) < 4:
            continue
        year_str, month_str = stem_parts[-2], stem_parts[-1]
        if not (year_str.isdigit() and month_str.isdigit()):
            continue
        year = int(year_str)
        month = int(month_str)
        if 1 <= month <= 12:
            months.add((year, month))
    return sorted(months)


def _month_span(*, start: tuple[int, int], end: tuple[int, int]) -> list[tuple[int, int]]:
    """Return inclusive (year, month) points between start and end."""
    months: list[tuple[int, int]] = []
    year, month = start
    while (year, month) <= end:
        months.append((year, month))
        month += 1
        if month > 12:
            month = 1
            year += 1
    return months


@app.get("/{asset_path:path}")
def next_export_assets(asset_path: str) -> FileResponse:
    """
    Serve static files from Next export output when available.
    Keeps legacy behavior by falling back to the old frontend index
    only for root route (handled above).
    """
    if not asset_path:
        raise HTTPException(status_code=404, detail="Not found")

    candidate = NEXT_OUT_DIR / asset_path
    if candidate.is_file():
        return FileResponse(str(candidate))

    # Support exported folder routes (e.g. /foo -> /foo/index.html).
    folder_index = candidate / "index.html"
    if folder_index.is_file():
        return FileResponse(str(folder_index))

    raise HTTPException(status_code=404, detail="Not found")


