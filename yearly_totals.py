from copy import copy
from pathlib import Path

from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook

from config import get_paths
from data import monthToExpenseColDict


def _month_column_letter(month: str | int) -> str:
    month_int = int(month)
    if month_int not in monthToExpenseColDict:
        raise ValueError(f"Invalid month for yearly total sync: {month!r}")
    return monthToExpenseColDict[month_int]


def _monthly_output_path(year: str, month: str | int) -> Path:
    paths = get_paths()
    month_padded = str(month).zfill(2)
    year_dir = paths.data_dir / str(year)
    yearly_path = year_dir / f"expenses_output_{year}_{month_padded}.xlsx"
    if yearly_path.exists():
        return yearly_path
    return paths.data_dir / f"expenses_output_{year}_{month_padded}.xlsx"


def _yearly_total_path(year: str) -> Path:
    paths = get_paths()
    year_dir = paths.data_dir / str(year)
    year_dir.mkdir(parents=True, exist_ok=True)
    return year_dir / f"expenses_output_{year}_total.xlsx"


def ensure_yearly_total_workbook(year: str) -> Path:
    paths = get_paths()
    total_path = _yearly_total_path(year)
    if total_path.exists():
        return total_path
    if not paths.template_expenses.exists():
        raise FileNotFoundError(
            f"Template workbook not found at {paths.template_expenses}. "
            "Cannot create yearly total workbook."
        )
    total_path.write_bytes(paths.template_expenses.read_bytes())
    return total_path


def sync_month_to_year_total(year: str, month: str | int) -> Path:
    source_month_path = _monthly_output_path(year, month)
    if not source_month_path.exists():
        raise FileNotFoundError(
            f"Monthly workbook not found at {source_month_path}. "
            "Cannot sync yearly total workbook."
        )

    total_path = ensure_yearly_total_workbook(year)
    col_letter = _month_column_letter(month)

    src_wb = load_workbook(source_month_path)
    dst_wb = load_workbook(total_path)
    try:
        src_ws = src_wb.active
        dst_ws = dst_wb.active
        max_rows = max(src_ws.max_row, dst_ws.max_row)

        for row in range(1, max_rows + 1):
            src_cell = src_ws[f"{col_letter}{row}"]
            dst_cell = dst_ws[f"{col_letter}{row}"]
            if isinstance(dst_cell, MergedCell):
                continue
            dst_cell.value = src_cell.value
            dst_cell.comment = copy(src_cell.comment) if src_cell.comment else None

        dst_wb.save(total_path)
    finally:
        src_wb.close()
        dst_wb.close()

    return total_path

