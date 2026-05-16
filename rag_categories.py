from __future__ import annotations

import json
import re
from dataclasses import dataclass
from difflib import SequenceMatcher
from typing import Dict, List, Tuple, Any, Optional

from openpyxl import load_workbook

from config import get_paths, template_category_rows


@dataclass
class HistoricalExpense:
    name: str
    amount: float
    category: str


_HISTORICAL_CACHE: Optional[List[HistoricalExpense]] = None
_HISTORICAL_CACHE_KEY: Optional[str] = None
_ALLOWED_SUBCATS_CACHE: Optional[List[str]] = None


def _normalize_for_fuzzy_name(s: str) -> str:
    return re.sub(r"[^א-תa-zA-Z0-9]", "", str(s or "").lower())


def _safe_float_from_string(s: str) -> Optional[float]:
    if s is None:
        return None
    if isinstance(s, (int, float)):
        return float(s)
    s2 = str(s).strip().replace("\u200f", "").replace("\u200e", "").replace("₪", "").replace(",", "")
    if not s2:
        return None
    try:
        return float(s2)
    except ValueError:
        m = re.search(r"[-+]?\d+(?:\.\d+)?", s2)
        if not m:
            return None
        try:
            return float(m.group(0))
        except ValueError:
            return None


def _get_history_workbooks() -> List[Tuple[str, float]]:
    """
    Return a list of (path, mtime) for recent output workbooks.
    """
    paths = get_paths()
    candidates = sorted(
        paths.data_dir.glob("**/expenses_output_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    picked = candidates[:3]
    if not picked and paths.current_expenses.exists():
        picked = [paths.current_expenses]
    return [(str(p), p.stat().st_mtime) for p in picked]


def _get_allowed_subcategories() -> List[str]:
    """
    Allowed expense sub-category labels (exact strings) from the template workbook.
    Used to filter historical examples so we only feed the model valid categories.
    """
    global _ALLOWED_SUBCATS_CACHE
    if _ALLOWED_SUBCATS_CACHE is not None:
        return _ALLOWED_SUBCATS_CACHE

    paths = get_paths()
    if not paths.template_expenses.exists():
        _ALLOWED_SUBCATS_CACHE = []
        return _ALLOWED_SUBCATS_CACHE

    wb = load_workbook(str(paths.template_expenses), data_only=True)
    ws = wb.active
    try:
        labels: List[str] = []
        seen = set()
        for r in template_category_rows():
            v = ws[f"B{r}"].value
            if v is None:
                continue
            s = str(v).strip()
            if not s or s in seen:
                continue
            seen.add(s)
            labels.append(s)
        _ALLOWED_SUBCATS_CACHE = labels
        return labels
    finally:
        wb.close()


def _load_historical_expenses() -> List[HistoricalExpense]:
    """
    Extract historical (expense name -> sub-category) examples from your filled Excel outputs.

    In handleExcel.fillCells(), each categorized expense adds a Comment to the month cell:
        "<expenseName> - <amount>"
    under the correct sub-category row (column B).
    """
    global _HISTORICAL_CACHE, _HISTORICAL_CACHE_KEY

    history = _get_history_workbooks()
    paths = get_paths()
    exp_dict_path = paths.data_dir / "expensesDict.txt"
    exp_dict_mtime = exp_dict_path.stat().st_mtime if exp_dict_path.exists() else 0.0
    if not history and not exp_dict_path.exists():
        print("[rag_categories] No output workbook found and no expensesDict.txt; skipping historical examples")
        return []

    cache_key = "|".join([*(f"{p}:{mt}" for p, mt in history), f"expensesDict:{exp_dict_mtime}"])
    if _HISTORICAL_CACHE is not None and _HISTORICAL_CACHE_KEY == cache_key:
        return _HISTORICAL_CACHE

    all_hist: List[HistoricalExpense] = []

    for workbook_path, _mtime in history:
        print(f"[rag_categories] Loading historical examples from: {workbook_path}")
        wb = load_workbook(str(workbook_path), data_only=True)
        ws = wb.active
        try:
            category_col = "B"
            month_cols = range(3, 15)  # C..N

            for row in template_category_rows():
                category = ws[f"{category_col}{row}"].value
                if not category:
                    continue
                category_str = str(category).strip()
                if not category_str:
                    continue

                for col in month_cols:
                    cell = ws.cell(row=row, column=col)
                    if cell.comment is None or not cell.comment.text:
                        continue
                    text = cell.comment.text
                    for line in text.splitlines():
                        line = line.strip()
                        if not line:
                            continue
                        if " - " not in line:
                            continue
                        name_part, amount_part = line.rsplit(" - ", 1)
                        name_part = str(name_part).strip()
                        amount = _safe_float_from_string(amount_part)
                        if not name_part or amount is None:
                            continue
                        all_hist.append(
                            HistoricalExpense(
                                name=name_part,
                                # Workbook "expenses" should never be negative; normalize to abs.
                                amount=float(abs(amount)),
                                category=category_str,
                            )
                        )
        finally:
            wb.close()

    if not all_hist and exp_dict_path.exists():
        # Fallback: parse the latest expensesDict.txt lines.
        # We only keep rows where the category is one of the template sub-categories.
        allowed_subcats = set(_get_allowed_subcategories())
        print(f"[rag_categories] No workbook comments found; falling back to: {exp_dict_path}")
        raw = exp_dict_path.read_text(encoding="utf-8", errors="replace")
        for line in raw.splitlines():
            line = line.strip()
            if not line or " - " not in line:
                continue
            parts = line.rsplit(" - ", 2)
            if len(parts) != 3:
                continue
            name, amount_str, category = parts
            category = str(category).strip()
            if category not in allowed_subcats:
                continue
            amount = _safe_float_from_string(amount_str)
            if amount is None:
                continue
            name = str(name).strip().replace("-", "~")
            all_hist.append(
                HistoricalExpense(name=name, amount=float(abs(amount)), category=category)
            )

    # De-dupe (keep the highest-absolute amount example just to avoid huge repeats)
    best_by_key: Dict[Tuple[str, str], HistoricalExpense] = {}
    for ex in all_hist:
        k = (ex.name, ex.category)
        cur = best_by_key.get(k)
        if cur is None or abs(ex.amount) > abs(cur.amount):
            best_by_key[k] = ex

    historical = list(best_by_key.values())
    print(f"[rag_categories] Loaded {len(historical)} historical expense examples")

    _HISTORICAL_CACHE = historical
    _HISTORICAL_CACHE_KEY = cache_key
    return historical


def get_category_corrections() -> Dict[str, str]:
    """
    Load user-defined corrections: expense name -> correct template category.
    Used to teach the model; add entries to data/category_corrections.json
    (format: {"שם הוצאה": "תת-קטגוריה מדויקת מהתבנית"}).
    """
    paths = get_paths()
    if not paths.category_corrections_file.exists():
        return {}
    try:
        with paths.category_corrections_file.open("r", encoding="utf-8") as f:
            data = json.load(f)
    except (json.JSONDecodeError, OSError):
        return {}
    return {k: v for k, v in (data or {}).items() if isinstance(k, str) and isinstance(v, str) and not k.startswith("_")}


def get_corrections_text_for_prompt() -> str:
    """Build prompt text from category_corrections.json so the model follows user corrections."""
    corrections = get_category_corrections()
    if not corrections:
        return ""
    lines: List[str] = []
    for name, category in corrections.items():
        name_norm = str(name).replace("-", "~")
        if name_norm != name:
            lines.append(f'"{name_norm}" -> "{category}"')
            lines.append(f'"{name}" -> "{category}"')
        else:
            lines.append(f'"{name}" -> "{category}"')
    return (
        "\n\nתיקונים מהמשתמש (חובה לכבד בהתאמה מדויקת – סווג את ההוצאות הבאות exactly לפי הקטגוריה שצוינה):\n"
        + "\n".join(lines)
    )


def get_examples_for_prompt(expenses_dict: Dict[str, Any], max_examples: int = 10) -> str:
    """
    Build a short, relevant examples string from your historical data.

    We retrieve examples that look similar to the current expense names (lightweight fuzzy match),
    then we add them as few-shot examples to help the model choose the right sub-category.
    """
    historical = _load_historical_expenses()
    if not historical:
        return ""

    # Precompute normalized historical names for cheap fuzzy matching.
    hist_norm: List[Tuple[str, HistoricalExpense]] = [
        (_normalize_for_fuzzy_name(ex.name), ex) for ex in historical
    ]
    hist_norm = [(n, ex) for (n, ex) in hist_norm if n]

    current_names = list(expenses_dict.keys())
    current_norm = {name: _normalize_for_fuzzy_name(name) for name in current_names}

    # Retrieve best match per current expense name.
    scored: List[Tuple[float, HistoricalExpense]] = []
    # Keep retrieval bounded so prompt building stays fast.
    hist_candidates = hist_norm[:2000] if len(hist_norm) > 2000 else hist_norm
    for name in current_names:
        cn = current_norm.get(name, "")
        if not cn:
            continue
        best_for_this = None
        best_score = 0.0
        for hn, ex in hist_candidates:
            if not hn:
                continue
            score = SequenceMatcher(None, cn, hn).ratio()
            if score > best_score:
                best_score = score
                best_for_this = ex
        if best_for_this is not None and best_score >= 0.55:
            scored.append((best_score, best_for_this))

    if not scored:
        return ""

    scored.sort(key=lambda x: x[0], reverse=True)
    chosen: List[HistoricalExpense] = []
    used_categories: set[str] = set()
    used_pairs: set[Tuple[str, str]] = set()
    for _score, ex in scored:
        key = (ex.name, ex.category)
        if key in used_pairs:
            continue
        # Prefer category diversity to avoid the prompt being dominated by one bucket.
        if ex.category in used_categories:
            continue
        chosen.append(ex)
        used_pairs.add(key)
        used_categories.add(ex.category)
        if len(chosen) >= max_examples:
            break

    if not chosen:
        return ""

    examples_text = "\n".join(
        f"{ex.name} - {ex.amount} - {ex.category}" for ex in chosen
    )
    out = (
        "\n\nהנה כמה דוגמאות מסיווגים קודמים (שם ההוצאה - סכום - תת קטגוריה):\n"
        + examples_text
    )
    corrections_text = get_corrections_text_for_prompt()
    if corrections_text:
        out += corrections_text
    return out

